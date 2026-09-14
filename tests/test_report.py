from pathlib import Path
from openpyxl import load_workbook

from excel2api.report import extract_path, write_sync_report
from excel2api.api import APIResult


def test_extract_path():
    assert extract_path({"data": {"id": 42}}, "data.id") == 42
    assert extract_path({"id": 1}, "missing", "x") == "x"


def test_write_sync_report(tmp_path: Path):
    output = tmp_path / "report.xlsx"
    results = [APIResult(2, "CREATE", True, 201, {"id": 99, "status": "created"})]
    write_sync_report(output, results, {"api_id": "id", "api_status": "status"})
    wb = load_workbook(output)
    ws = wb["Sync Results"]
    assert ws["G2"].value == 99
    assert ws["H2"].value == "created"
