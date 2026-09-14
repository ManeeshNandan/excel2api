from pathlib import Path
from openpyxl import Workbook
from excel2api.reader import list_sheets, read_file
from excel2api.converter import convert_file


def test_read_specific_sheet(tmp_path: Path):
    path = tmp_path / 'data.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Patients'
    ws.append(['Name']); ws.append(['John'])
    ws2 = wb.create_sheet('Doctors')
    ws2.append(['Name']); ws2.append(['Rao'])
    wb.save(path)
    assert list_sheets(path) == ['Patients', 'Doctors']
    assert read_file(path, sheet_name='Doctors')[0]['Name'] == 'Rao'


def test_convert_specific_sheet(tmp_path: Path):
    path = tmp_path / 'data.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Patients'
    ws.append(['Name']); ws.append(['John']); wb.create_sheet('Doctors').append(['Name'])
    wb.save(path)
    schema = tmp_path / 'schema.yaml'
    schema.write_text('fields:\n  name:\n    column: Name\n    required: true\n')
    records, errors = convert_file(path, schema, sheet_name='Patients')
    assert records == [{'name': 'John'}]
    assert errors == []
